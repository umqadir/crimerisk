#!/usr/bin/env python3
import csv
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[6]
PACKET_DIR = Path(__file__).resolve().parent
DATA_DIR = REPO_ROOT / "data" / "FDLE-FIBRS-2024"
WORKBOOK_PATH = DATA_DIR / "raw" / "FIBRS_Offense_Detail_2021-2026.xlsx"
SPOTCHECK_PATH = PACKET_DIR / "official_fibrs_spotcheck_2024.csv"
NORMALIZED_PATH = DATA_DIR / "parsed" / "fdle_fibrs_2024_offense_period_extract.csv"
SCHEMA_PATH = PACKET_DIR / "fdle_fibrs_2024_workbook_schema.json"
CONTRACT_PATH = PACKET_DIR / "fdle_fibrs_2024_ingestion_contract.json"
RECONCILIATION_PATH = PACKET_DIR / "official_fibrs_spotcheck_2024_reconciled.csv"

STATE_ABBR = "FL"
YEAR = 2024
SHEET_NAME = "FIBRS Offense 2024"
SOURCE_URL = (
    "https://www.fdle.state.fl.us/getContentAsset/43da39b3-e350-4c0a-b5f2-2a9a98b6132e/"
    "73aabf56-e6e5-4330-95a3-5f2a270a1d2b/FIBRS_Offense_Detail_2021-2026.xlsx?language=en"
)
SOURCE_LANDING_PAGE = "https://www.fdle.state.fl.us/cjab/fibrs"

DETAIL_SECTION_STARTS = {4: "violent", 52: "property", 172: "other"}
DETAIL_BLOCK_WIDTH = 12
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ALIAS_ONLY_TOKENS = {
    "police",
    "sheriff",
    "dept",
    "office",
    "resident",
    "agency",
    "field",
    "campus",
}

OFFENSE_NAME_CORRECTIONS = {
    "Prostitutition": "Prostitution",
}


def clean_string(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def parse_count(value):
    if value in (None, "", "--"):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    value = str(value).strip()
    if not value or value == "--":
        return None
    number = float(value)
    return int(number) if number.is_integer() else number


def parse_rate(value):
    if value in (None, "", "--"):
        return None
    return float(value)


def reported_month_indexes(values):
    return [
        index
        for index, value in enumerate(values, start=1)
        if value not in (None, "", "--")
    ]


def quarters_from_months(month_indexes):
    if not month_indexes:
        return None
    return len({int((month - 1) // 3) + 1 for month in month_indexes})


def slugify(value):
    value = OFFENSE_NAME_CORRECTIONS.get(value, value)
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def normalize_agency_name(value):
    if not value:
        return ""
    value = value.lower()
    replacements = [
        ("&", " and "),
        ("department of law enforcement", "fdle"),
        ("dept. of law enforcement", "fdle"),
        ("department", "dept"),
        ("dept.", "dept"),
        ("police department", "police"),
        ("police dept.", "police"),
        ("police dept", "police"),
        (" sheriff's office", " sheriff"),
        (" sheriffs office", " sheriff"),
        (" sheriff office", " sheriff"),
        (" sheriff's", " sheriff"),
        (" sheriffs", " sheriff"),
        ("univ.", "university"),
        ("univ ", "university "),
        (" pd ", " police "),
        (" pd", " police"),
        (" ra", " resident agency"),
        ("tboc", " field office"),
        (":", " "),
        ("-", " "),
        ("/", " "),
        (",", " "),
    ]
    for old, new in replacements:
        value = value.replace(old, new)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def tokenize(value):
    return set(normalize_agency_name(value).split())


def build_detail_blocks(ws):
    row2 = [ws.cell(2, column).value for column in range(1, ws.max_column + 1)]
    blocks = []
    section = None
    for column in range(4, ws.max_column + 1):
        if clean_string(row2[column - 1]) is None:
            continue
        if column in DETAIL_SECTION_STARTS:
            section = DETAIL_SECTION_STARTS[column]
        offense_name_source = clean_string(row2[column - 1])
        offense_name = OFFENSE_NAME_CORRECTIONS.get(offense_name_source, offense_name_source)
        blocks.append(
            {
                "section": section,
                "start_column": column,
                "offense_name_source": offense_name_source,
                "offense_name": offense_name,
                "offense_slug": slugify(offense_name_source),
            }
        )
    return blocks


def read_sheet():
    wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]
    detail_blocks = build_detail_blocks(ws)

    agencies = []
    extract_rows = []
    months_reported_counts = Counter()

    for row_number, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        county = clean_string(row[0])
        agency_name = clean_string(row[1])
        if county is None and agency_name is None:
            continue
        population_estimate = parse_count(row[2])

        agency_record = {
            "row_number": row_number,
            "county": county,
            "agency_name": agency_name,
            "agency_name_normalized": normalize_agency_name(agency_name),
            "population_estimate": population_estimate,
            "months_reported": None,
            "quarters_reported": None,
            "summary_counts": {"violent": {"ytd": 0}, "property": {"ytd": 0}, "other": {"ytd": 0}},
        }

        agency_month_indexes = set()
        for block in detail_blocks:
            start = block["start_column"] - 1
            raw_month_values = list(row[start : start + DETAIL_BLOCK_WIDTH])
            month_indexes = reported_month_indexes(raw_month_values)
            agency_month_indexes.update(month_indexes)
            month_counts = [parse_count(value) for value in raw_month_values]
            months_reported = len(month_indexes)
            quarters_reported = quarters_from_months(month_indexes)
            ytd = sum(count for count in month_counts if count is not None) if month_indexes else None
            if ytd is not None and block["section"] in agency_record["summary_counts"]:
                agency_record["summary_counts"][block["section"]]["ytd"] += ytd

            for month_index, count in enumerate(month_counts, start=1):
                extract_rows.append(
                    {
                        "state_abbr": STATE_ABBR,
                        "year": YEAR,
                        "source_sheet": SHEET_NAME,
                        "source_row_number": row_number,
                        "county": county,
                        "agency_name": agency_name,
                        "agency_name_normalized": agency_record["agency_name_normalized"],
                        "population_estimate": population_estimate,
                        "quarters_reported": quarters_reported,
                        "months_reported": months_reported,
                        "offense_level": "detail_offense",
                        "offense_bucket": block["section"],
                        "offense_name": block["offense_name"],
                        "offense_name_source": block["offense_name_source"],
                        "offense_slug": block["offense_slug"],
                        "period": f"M{month_index:02d}",
                        "quarter": int((month_index - 1) // 3) + 1,
                        "month": month_index,
                        "month_name": MONTH_NAMES[month_index - 1],
                        "incident_count": count,
                        "annual_rate": None,
                    }
                )

            extract_rows.append(
                {
                    "state_abbr": STATE_ABBR,
                    "year": YEAR,
                    "source_sheet": SHEET_NAME,
                    "source_row_number": row_number,
                    "county": county,
                    "agency_name": agency_name,
                    "agency_name_normalized": agency_record["agency_name_normalized"],
                    "population_estimate": population_estimate,
                    "quarters_reported": quarters_reported,
                    "months_reported": months_reported,
                    "offense_level": "detail_offense",
                    "offense_bucket": block["section"],
                    "offense_name": block["offense_name"],
                    "offense_name_source": block["offense_name_source"],
                    "offense_slug": block["offense_slug"],
                    "period": "YTD",
                    "quarter": None,
                    "month": None,
                    "month_name": None,
                    "incident_count": ytd,
                    "annual_rate": None,
                }
            )

        agency_record["months_reported"] = len(agency_month_indexes)
        agency_record["quarters_reported"] = quarters_from_months(agency_month_indexes)
        months_reported_counts[agency_record["months_reported"]] += 1
        agencies.append(agency_record)

    wb.close()
    return agencies, detail_blocks, extract_rows, months_reported_counts


def write_csv(path, rows, fieldnames):
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def choose_spotcheck_match(spotcheck_row, agencies):
    source_name = spotcheck_row["agency_name"]
    county = clean_string(spotcheck_row.get("fdle_county"))
    normalized_source = normalize_agency_name(source_name)
    source_tokens = tokenize(source_name)

    exact_county = [
        agency
        for agency in agencies
        if agency["county"] == county and agency["agency_name_normalized"] == normalized_source
    ]
    if exact_county:
        return exact_county[0], "county_normalized_name"

    exact_any = [agency for agency in agencies if agency["agency_name_normalized"] == normalized_source]
    if len(exact_any) == 1:
        return exact_any[0], "statewide_normalized_name"

    county_pool = [agency for agency in agencies if agency["county"] == county] if county else agencies

    if source_name.startswith("Department of Law Enforcement:"):
        city_fragment = source_name.split(",")[-1].strip() if "," in source_name else ""
        fdle_candidates = [
            agency
            for agency in county_pool
            if agency["agency_name"].startswith("FDLE - ")
            and (not city_fragment or city_fragment.lower() in agency["agency_name"].lower())
        ]
        if len(fdle_candidates) == 1:
            return fdle_candidates[0], "county_fdle_city_alias"

    containment_matches = []
    for agency in county_pool:
        candidate_tokens = tokenize(agency["agency_name"])
        overlap = source_tokens & candidate_tokens
        if not source_tokens:
            continue
        if source_tokens <= candidate_tokens:
            extra_tokens = candidate_tokens - source_tokens
        elif candidate_tokens <= source_tokens:
            extra_tokens = source_tokens - candidate_tokens
        else:
            continue
        if extra_tokens and not extra_tokens <= ALIAS_ONLY_TOKENS:
            continue
        containment_matches.append((len(overlap), agency))
    if containment_matches:
        containment_matches.sort(key=lambda item: (-item[0], item[1]["agency_name"]))
        return containment_matches[0][1], "county_token_containment_alias"

    return None, None


def reconcile_spotcheck(agencies):
    agency_lookup = {(agency["county"], agency["agency_name"]): agency for agency in agencies}
    with SPOTCHECK_PATH.open(newline="") as handle:
        source_rows = list(csv.DictReader(handle))

    reconciled_rows = []
    status_counts = Counter()

    for row in source_rows:
        matched_agency, match_method = choose_spotcheck_match(row, agencies)
        existing_key = (clean_string(row.get("fdle_county")), clean_string(row.get("fdle_agency_name")))
        existing_agency = agency_lookup.get(existing_key) if existing_key[1] else None

        if matched_agency is None:
            reconciliation_status = "not_found_in_workbook"
            status_counts[reconciliation_status] += 1
            matched_summary = {"violent": None, "property": None, "other": None}
            matched_county = None
            matched_name = None
            matched_quarters = None
        else:
            matched_summary = matched_agency["summary_counts"]
            matched_county = matched_agency["county"]
            matched_name = matched_agency["agency_name"]
            matched_quarters = matched_agency["quarters_reported"]
            if existing_agency and existing_agency["agency_name"] == matched_name and existing_agency["county"] == matched_county:
                reconciliation_status = "matched_same_workbook_row"
            elif existing_agency:
                reconciliation_status = "matched_different_workbook_row"
            else:
                reconciliation_status = "matched_new_from_normalized_extract"
            status_counts[reconciliation_status] += 1

        existing_counts_match = None
        if existing_agency:
            existing_counts_match = all(
                parse_count(row[field]) == existing_agency["summary_counts"][bucket]["ytd"]
                for field, bucket in [
                    ("violent_ytd", "violent"),
                    ("property_ytd", "property"),
                    ("other_ytd", "other"),
                ]
            ) and parse_count(row["fdle_quarters_reported"]) == existing_agency["quarters_reported"]

        matched_counts_match_existing = None
        if matched_agency and existing_agency:
            matched_counts_match_existing = all(
                matched_summary[bucket]["ytd"] == existing_agency["summary_counts"][bucket]["ytd"]
                for bucket in ("violent", "property", "other")
            ) and matched_agency["quarters_reported"] == existing_agency["quarters_reported"]

        note_parts = []
        if matched_agency is None:
            note_parts.append("No 2024 workbook row located for this spotcheck agency.")
        elif reconciliation_status == "matched_new_from_normalized_extract":
            note_parts.append("Recovered a workbook match not captured in the original spotcheck.")
        if match_method == "county_token_containment_alias":
            note_parts.append("Alias match relies on county plus token containment rather than exact normalized name.")
        if match_method == "county_fdle_city_alias":
            note_parts.append("Alias match relies on county plus FDLE regional office city token.")
        if existing_counts_match is False:
            note_parts.append("Original spotcheck counts do not match the workbook row named in the spotcheck.")

        reconciled_rows.append(
            {
                "ori9": row["ori9"],
                "agency_name": row["agency_name"],
                "agency_type_name": row["agency_type_name"],
                "official_nibrs_start_date": row["official_nibrs_start_date"],
                "original_spotcheck_status": row["spotcheck_status"],
                "reconciliation_status": reconciliation_status,
                "match_method": match_method,
                "matched_county": matched_county,
                "matched_agency_name": matched_name,
                "matched_quarters_reported": matched_quarters,
                "matched_violent_ytd": matched_summary["violent"]["ytd"] if matched_agency else None,
                "matched_property_ytd": matched_summary["property"]["ytd"] if matched_agency else None,
                "matched_other_ytd": matched_summary["other"]["ytd"] if matched_agency else None,
                "existing_fdle_county": row["fdle_county"] or None,
                "existing_fdle_agency_name": row["fdle_agency_name"] or None,
                "existing_fdle_quarters_reported": parse_count(row["fdle_quarters_reported"]),
                "existing_violent_ytd": parse_count(row["violent_ytd"]),
                "existing_property_ytd": parse_count(row["property_ytd"]),
                "existing_other_ytd": parse_count(row["other_ytd"]),
                "existing_spotcheck_counts_match_workbook": existing_counts_match,
                "matched_counts_equal_existing_row": matched_counts_match_existing,
                "notes": " ".join(note_parts) or None,
            }
        )

    return reconciled_rows, status_counts


def build_schema(detail_blocks, extract_rows, agencies, months_reported_counts):
    return {
        "state_abbr": STATE_ABBR,
        "source_workbook_path": str(WORKBOOK_PATH),
        "source_workbook_url": SOURCE_URL,
        "source_landing_page": SOURCE_LANDING_PAGE,
        "worksheet": SHEET_NAME,
        "header_rows": [1, 2, 3],
        "first_data_row": 4,
        "agency_row_count": len(agencies),
        "extract_row_count": len(extract_rows),
        "identifier_columns": [
            {"column_number": 1, "column_name": "County", "normalized_field": "county"},
            {"column_number": 2, "column_name": "Agency Name", "normalized_field": "agency_name"},
            {
                "column_number": 3,
                "column_name": "2024 Population Estimate",
                "normalized_field": "population_estimate",
            },
        ],
        "detail_offense_blocks": detail_blocks,
        "extract_columns": [
            "state_abbr",
            "year",
            "source_sheet",
            "source_row_number",
            "county",
            "agency_name",
            "agency_name_normalized",
            "population_estimate",
            "quarters_reported",
            "months_reported",
            "offense_level",
            "offense_bucket",
            "offense_name",
            "offense_name_source",
            "offense_slug",
            "period",
            "quarter",
            "month",
            "month_name",
            "incident_count",
            "annual_rate",
        ],
        "known_schema_issues": [
            "Workbook does not expose ORI values, so agency joins must use county and agency-name crosswalk logic.",
            "The source label 'Prostitutition' is misspelled in the workbook; the extract preserves the source label and normalizes the slug to 'prostitution'.",
            "The refreshed workbook changed from quarter/YTD columns to monthly detail columns; YTD rows in the extract are computed from monthly cells.",
        ],
        "months_reported_distribution": dict(sorted(months_reported_counts.items())),
        "quarters_reported_distribution": dict(
            sorted(Counter(agency["quarters_reported"] for agency in agencies).items())
        ),
        "agencies_with_zero_population_estimate": sum(1 for agency in agencies if agency["population_estimate"] == 0),
    }


def build_contract(schema, reconciliation_rows, reconciliation_status_counts):
    matched_rows = [
        row
        for row in reconciliation_rows
        if row["reconciliation_status"] != "not_found_in_workbook"
    ]
    return {
        "contract_version": "2026-07-02",
        "state_abbr": STATE_ABBR,
        "year": YEAR,
        "source": {
            "official_source_type": "FDLE FIBRS workbook",
            "official_source_url": SOURCE_URL,
            "landing_page_url": SOURCE_LANDING_PAGE,
            "local_workbook_path": str(WORKBOOK_PATH),
            "selected_worksheet": SHEET_NAME,
        },
        "selection_rules": {
            "use_only_selected_worksheet_for_2024": True,
            "skip_header_rows": [1, 2, 3],
            "treat_row_4_and_below_as_agency_rows": True,
        },
        "normalization_rules": {
            "county_field": "trimmed source County cell",
            "agency_name_field": "trimmed source Agency Name cell",
            "agency_name_normalized": "lowercased punctuation-stripped alias key for county-aware joins",
            "population_estimate": "numeric source column 3",
            "months_reported": "count of non-null monthly cells by agency/offense; zero counts still count as reported months",
            "quarters_reported": "derived from the reported monthly cells for compatibility with older state-publication consumers",
            "count_null_token": "--",
            "count_columns": "monthly cells converted to nullable integers; YTD rows are computed as the sum of monthly counts",
        },
        "staging_output": {
            "path": str(NORMALIZED_PATH),
            "grain": "one row per agency x offense/group x period",
            "period_values": ["M01", "M02", "M03", "M04", "M05", "M06", "M07", "M08", "M09", "M10", "M11", "M12", "YTD"],
            "offense_levels": ["detail_offense"],
            "row_count": schema["extract_row_count"],
            "required_columns": schema["extract_columns"],
        },
        "join_contract": {
            "preferred_join_keys": [
                "county",
                "agency_name_normalized",
            ],
            "join_limitations": [
                "No ORI appears in the workbook.",
                "Alias resolution is required for some college and FDLE regional-office rows.",
                "Many validation rows remain genuine workbook absences rather than naming misses.",
            ],
        },
        "known_matching_issues": [
            {
                "issue": "county_sheriff_absences",
                "examples": [
                    "Hamilton County Sheriff's Office",
                    "Highlands County Sheriff's Office",
                    "Okaloosa County Sheriff's Office",
                    "Union County Sheriff's Office",
                ],
                "note": "These roster rows were not found in the refreshed 2024 FDLE workbook.",
            },
            {
                "issue": "city_absences",
                "examples": [
                    "Plant City Police Department",
                    "Dade City Police Department",
                    "Port Richey Police Department",
                    "Venice Police Department",
                    "Southwest Ranches Police Department",
                ],
                "note": "These spotcheck agencies were not found as refreshed 2024 workbook rows.",
            },
            {
                "issue": "alias_only_matches",
                "examples": [
                    "Tallahassee Community College -> Tallahassee Community College PD",
                    "Department of Law Enforcement: Lee County, Fort Myers -> FDLE - Fort Myers RA",
                    "Department of Law Enforcement: Leon County, Tallahassee -> FDLE - Tallahassee",
                ],
                "note": "These require county-aware alias logic; exact normalized-name joins are insufficient.",
            },
            {
                "issue": "campus_name_mismatch",
                "examples": [
                    "University of South Florida: Tampa",
                ],
                "note": "The workbook only shows 'Univ of South Florida PD - St Pete Campus' and does not expose a Tampa campus row.",
            },
        ],
        "spotcheck_reconciliation": {
            "source_path": str(SPOTCHECK_PATH),
            "output_path": str(RECONCILIATION_PATH),
            "row_count": len(reconciliation_rows),
            "matched_row_count": len(matched_rows),
            "status_counts": dict(sorted(reconciliation_status_counts.items())),
        },
    }


def main():
    agencies, detail_blocks, extract_rows, months_reported_counts = read_sheet()

    write_csv(
        NORMALIZED_PATH,
        extract_rows,
        [
            "state_abbr",
            "year",
            "source_sheet",
            "source_row_number",
            "county",
            "agency_name",
            "agency_name_normalized",
            "population_estimate",
            "quarters_reported",
            "months_reported",
            "offense_level",
            "offense_bucket",
            "offense_name",
            "offense_name_source",
            "offense_slug",
            "period",
            "quarter",
            "month",
            "month_name",
            "incident_count",
            "annual_rate",
        ],
    )

    reconciliation_rows, reconciliation_status_counts = reconcile_spotcheck(agencies)
    write_csv(
        RECONCILIATION_PATH,
        reconciliation_rows,
        [
            "ori9",
            "agency_name",
            "agency_type_name",
            "official_nibrs_start_date",
            "original_spotcheck_status",
            "reconciliation_status",
            "match_method",
            "matched_county",
            "matched_agency_name",
            "matched_quarters_reported",
            "matched_violent_ytd",
            "matched_property_ytd",
            "matched_other_ytd",
            "existing_fdle_county",
            "existing_fdle_agency_name",
            "existing_fdle_quarters_reported",
            "existing_violent_ytd",
            "existing_property_ytd",
            "existing_other_ytd",
            "existing_spotcheck_counts_match_workbook",
            "matched_counts_equal_existing_row",
            "notes",
        ],
    )

    schema = build_schema(
        detail_blocks=detail_blocks,
        extract_rows=extract_rows,
        agencies=agencies,
        months_reported_counts=months_reported_counts,
    )
    contract = build_contract(schema, reconciliation_rows, reconciliation_status_counts)

    SCHEMA_PATH.write_text(json.dumps(schema, indent=2) + "\n")
    CONTRACT_PATH.write_text(json.dumps(contract, indent=2) + "\n")

    print(f"Wrote normalized extract: {NORMALIZED_PATH}")
    print(f"Wrote schema summary: {SCHEMA_PATH}")
    print(f"Wrote ingestion contract: {CONTRACT_PATH}")
    print(f"Wrote spotcheck reconciliation: {RECONCILIATION_PATH}")


if __name__ == "__main__":
    main()
